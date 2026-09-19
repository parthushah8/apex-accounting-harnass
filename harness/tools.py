from __future__ import annotations

import csv
import io
import subprocess
import sys
import time
import traceback
from pathlib import Path

TOOL_SCHEMA = [
    {
        "type": "function",
        "function": {
            "name": "list_dir",
            "description": "List files and directories under path (workspace-relative).",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "default": "."},
                },
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "read_file",
            "description": "Read a text/csv/txt file. For xlsx use inspect_xlsx. For pdf use read_pdf.",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "offset": {"type": "integer", "default": 0},
                    "max_chars": {"type": "integer", "default": 12000},
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "inspect_xlsx",
            "description": "Sheet names and first rows of an xlsx workbook, or header + sample rows of a csv.",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "max_rows": {"type": "integer", "default": 12},
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "read_xlsx",
            "description": "Read an xlsx sheet or csv file as a text table (csv: sheet param ignored).",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "sheet": {"type": "string"},
                    "skip": {"type": "integer", "default": 0},
                    "max_rows": {"type": "integer", "default": 40},
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "read_pdf",
            "description": "Extract text from a PDF.",
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string"},
                    "max_chars": {"type": "integer", "default": 12000},
                },
                "required": ["path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "run_python",
            "description": "Execute Python in the workspace. pandas, openpyxl, numpy, pypdf already available. Print results.",
            "parameters": {
                "type": "object",
                "properties": {"code": {"type": "string"}},
                "required": ["code"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "submit_answer",
            "description": "Submit the final console answer. Ends the run. Call once.",
            "parameters": {
                "type": "object",
                "properties": {"answer": {"type": "string"}},
                "required": ["answer"],
            },
        },
    },
]


class ToolError(Exception):
    pass


class Toolbelt:
    def __init__(self, workspace: Path):
        self.workspace = workspace.resolve()
        self.touched: list[dict] = []
        self.answer: str | None = None

    def resolve(self, path: str) -> Path:
        raw = (path or ".").replace("\\", "/").lstrip("/")
        if raw.startswith("/") or ".." in Path(raw).parts:
            raise ToolError(f"path escape blocked: {path}")
        target = self.workspace / (raw if raw != "." else "")
        # Do not Path.resolve() — workspace files are symlinks into apex-accounting.
        return target

    def touch(self, path: Path, action: str) -> str:
        if path.name.startswith(".__"):
            return path.name
        try:
            rel = "." if path == self.workspace else str(path.relative_to(self.workspace))
        except ValueError:
            rel = path.name
        rel = rel.replace("\\", "/")
        self.touched.append({"path": rel, "action": action, "t": time.time()})
        return rel

    def list_dir(self, path: str = ".") -> str:
        target = self.resolve(path)
        if not target.exists():
            raise ToolError(f"not found: {path}")
        if target.is_file():
            rel = self.touch(target, "stat")
            return f"{rel}  {target.stat().st_size} bytes"
        self.touch(target, "list")
        lines = []
        for child in sorted(target.iterdir(), key=lambda p: (not p.is_dir(), p.name.lower())):
            kind = "dir" if child.is_dir() else "file"
            size = "" if child.is_dir() else f"{child.stat().st_size:>10}"
            lines.append(f"{kind:4} {size}  {child.relative_to(self.workspace)}")
        return "\n".join(lines) or "(empty)"

    def read_file(self, path: str, offset: int = 0, max_chars: int = 12000) -> str:
        target = self.resolve(path)
        if not target.is_file():
            raise ToolError(f"not a file: {path}")
        suffix = target.suffix.lower()
        if suffix in {".xlsx", ".xls"}:
            raise ToolError("xlsx — use inspect_xlsx / read_xlsx")
        if suffix == ".pdf":
            raise ToolError("pdf — use read_pdf")
        text = target.read_text(errors="replace")
        self.touch(target, "read")
        chunk = text[offset : offset + max_chars]
        more = max(0, len(text) - offset - len(chunk))
        header = f"# {target.name}  chars={len(text)} offset={offset}\n"
        if more:
            header += f"# truncated, {more} chars remain — raise offset\n"
        return header + chunk

    def _csv_rows(self, target: Path) -> list[list[str]]:
        with target.open(newline="", encoding="utf-8-sig", errors="replace") as f:
            return list(csv.reader(f))

    def inspect_xlsx(self, path: str, max_rows: int = 12) -> str:
        target = self.resolve(path)
        if target.suffix.lower() == ".csv":
            self.touch(target, "inspect")
            rows = self._csv_rows(target)
            out = [f"# {target.name}  format=csv  rows={len(rows)}"]
            out.append(f"\n=== (csv)  rows={len(rows)} cols={max((len(r) for r in rows), default=0)} ===")
            for i, row in enumerate(rows[:max_rows], 1):
                out.append(f"{i:>3} | " + " | ".join(row))
            if len(rows) > max_rows:
                out.append(f"# {len(rows) - max_rows} rows remain")
            return "\n".join(out)

        import openpyxl

        wb = openpyxl.load_workbook(target, data_only=True)
        self.touch(target, "inspect")
        out = [f"# {target.name}  sheets={wb.sheetnames}"]
        for name in wb.sheetnames:
            ws = wb[name]
            out.append(f"\n=== {name}  rows={ws.max_row} cols={ws.max_column} ===")
            for i, row in enumerate(ws.iter_rows(max_row=min(max_rows, ws.max_row), values_only=True), 1):
                out.append(f"{i:>3} | " + " | ".join("" if v is None else str(v) for v in row))
        return "\n".join(out)

    def read_xlsx(self, path: str, sheet: str | None = None, skip: int = 0, max_rows: int = 40) -> str:
        target = self.resolve(path)
        if target.suffix.lower() == ".csv":
            rows = self._csv_rows(target)
            self.touch(target, "read")
            chunk = rows[skip : skip + max_rows]
            out = [f"# {target.name} / (csv)  total_rows={len(rows)} skip={skip}"]
            for i, row in enumerate(chunk, skip + 1):
                out.append(f"{i:>3} | " + " | ".join(row))
            if skip + max_rows < len(rows):
                out.append(f"# {len(rows) - skip - max_rows} rows remain")
            return "\n".join(out)

        import openpyxl

        wb = openpyxl.load_workbook(target, data_only=True)
        name = sheet or wb.sheetnames[0]
        if name not in wb.sheetnames:
            raise ToolError(f"no sheet {name!r}. have {wb.sheetnames}")
        ws = wb[name]
        self.touch(target, "read")
        rows = list(ws.iter_rows(values_only=True))
        chunk = rows[skip : skip + max_rows]
        out = [f"# {target.name} / {name}  total_rows={len(rows)} skip={skip}"]
        for i, row in enumerate(chunk, skip + 1):
            out.append(f"{i:>3} | " + " | ".join("" if v is None else str(v) for v in row))
        if skip + max_rows < len(rows):
            out.append(f"# {len(rows) - skip - max_rows} rows remain")
        return "\n".join(out)

    def read_pdf(self, path: str, max_chars: int = 12000) -> str:
        from pypdf import PdfReader

        target = self.resolve(path)
        reader = PdfReader(str(target))
        self.touch(target, "read")
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        chunk = text[:max_chars]
        more = max(0, len(text) - max_chars)
        header = f"# {target.name}  pages={len(reader.pages)} chars={len(text)}\n"
        if more:
            header += f"# truncated, {more} chars remain\n"
        return header + chunk

    def run_python(self, code: str) -> str:
        cleaned = code.strip()
        if cleaned.startswith("```"):
            cleaned = cleaned.strip("`")
            cleaned = cleaned.split("\n", 1)[-1]
        script = self.workspace / ".__agent_exec.py"
        script.write_text(cleaned)
        self.touch(script, "exec")
        t0 = time.time()
        try:
            proc = subprocess.run(
                [sys.executable, "-u", str(script)],
                cwd=self.workspace,
                capture_output=True,
                text=True,
                timeout=45,
            )
        except subprocess.TimeoutExpired as exc:
            raise ToolError("python timed out at 45s") from exc
        finally:
            if script.exists():
                script.unlink()
        ms = int((time.time() - t0) * 1000)
        out = (proc.stdout or "") + (("\n[stderr]\n" + proc.stderr) if proc.stderr else "")
        if proc.returncode != 0:
            raise ToolError(f"exit {proc.returncode} ({ms}ms)\n{out or '(no output)'}")
        return f"# ok {ms}ms\n{out or '(no stdout)'}"

    def submit_answer(self, answer: str) -> str:
        self.answer = answer
        return "submitted"

    def dispatch(self, name: str, args: dict) -> tuple[str, str | None]:
        fn = getattr(self, name, None)
        if fn is None:
            raise ToolError(f"unknown tool {name}")
        result = fn(**args)
        return result, self.answer


def preview(text: str, limit: int = 4000) -> str:
    if len(text) <= limit:
        return text
    return text[:limit] + f"\n… [{len(text) - limit} more chars]"
